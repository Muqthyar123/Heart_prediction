import io
import json
import os
import uuid
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Font, PatternFill
import pandas as pd
from app.extensions import db
from app.models.bulk import BulkJob, BulkPredictionResult
from app.services.model_loader import ModelManager
from app.utils.errors import (
    BadRequestError,
    ForbiddenError,
    NotFoundError,
    PayloadTooLargeError,
    ValidationError,
)
from app.utils.file_handling import save_secure_file
from app.utils.validation import REQUIRED_CLINICAL_FEATURES, validate_clinical_inputs


class BulkPredictionService:
    @staticmethod
    def read_spreadsheet(file_path: str) -> pd.DataFrame:
        ext = file_path.rsplit(".", 1)[-1].lower()
        try:
            if ext == "csv":
                df = pd.read_csv(file_path)
            elif ext in ("xlsx", "xls"):
                df = pd.read_excel(file_path)
            else:
                raise BadRequestError(f"Unsupported file format: .{ext}")
        except Exception as e:
            raise BadRequestError(f"Failed to read file: {str(e)}")

        if df.empty:
            raise BadRequestError("Uploaded file is empty.")

        # Normalize column names: strip whitespace, lowercase
        df.columns = [str(col).strip().lower() for col in df.columns]
        return df

    @staticmethod
    def validate_file_content(df: pd.DataFrame, max_rows: int) -> dict:
        total_rows = len(df)
        if total_rows > max_rows:
            raise PayloadTooLargeError(
                f"File contains {total_rows} rows, exceeding the maximum allowed limit of {max_rows} rows."
            )

        # Check required columns
        present_cols = set(df.columns)
        missing_cols = [c for c in REQUIRED_CLINICAL_FEATURES if c not in present_cols]
        extra_cols = [c for c in present_cols if c not in REQUIRED_CLINICAL_FEATURES and c != "patient_id"]

        if missing_cols:
            raise ValidationError(
                message=f"Missing required columns in uploaded file: {', '.join(missing_cols)}",
                details=[{"column": col, "error": "Missing column"} for col in missing_cols],
            )

        row_errors = []
        valid_count = 0

        for idx, row in df.iterrows():
            row_num = idx + 2  # Excel 1-based header offset
            row_data = row.to_dict()
            _, errs = validate_clinical_inputs(row_data)
            if errs:
                row_errors.append({
                    "row": row_num,
                    "errors": errs,
                })
            else:
                valid_count += 1

        return {
            "total_rows": total_rows,
            "valid_rows": valid_count,
            "invalid_rows": len(row_errors),
            "missing_columns": missing_cols,
            "detected_columns": list(df.columns),
            "row_errors": row_errors,
            "is_valid": len(row_errors) == 0,
        }

    @staticmethod
    def process_bulk_prediction(
        file_storage,
        upload_folder: str,
        user_id: str = None,
        max_rows: int = 5000,
        sync_limit: int = 500,
        model_name: str = "hybrid_stacking",
    ) -> dict:
        bulk_dir = os.path.join(upload_folder, "bulk")
        saved_name, full_path, _ = save_secure_file(
            file_storage,
            target_dir=bulk_dir,
            allowed_extensions={"csv", "xlsx", "xls"},
        )

        df = BulkPredictionService.read_spreadsheet(full_path)
        val_summary = BulkPredictionService.validate_file_content(df, max_rows)

        if not val_summary["is_valid"]:
            # Clean up uploaded file
            try:
                os.remove(full_path)
            except OSError:
                pass
            raise ValidationError(
                message="File content validation failed. Please correct the errors and re-upload.",
                details=val_summary["row_errors"][:50],  # Return first 50 row errors
            )

        model_mgr = ModelManager.get_instance()
        batch_results = model_mgr.predict_batch(df, model_name=model_name)

        pos_count = sum(1 for r in batch_results if r["prediction_class"] == 1)
        neg_count = sum(1 for r in batch_results if r["prediction_class"] == 0)

        # Create output dataframe with appended predictions
        df_out = df.copy()
        df_out["predicted_class"] = [r["prediction_class"] for r in batch_results]
        df_out["predicted_label"] = [r["prediction_label"] for r in batch_results]
        df_out["predicted_probability"] = [
            f"{r['probability']:.4f}" if r["probability"] is not None else "N/A"
            for r in batch_results
        ]

        # Save result file
        result_filename = f"result_{saved_name}.xlsx"
        result_path = os.path.join(bulk_dir, result_filename)
        df_out.to_excel(result_path, index=False)

        # Persist Job record
        job = BulkJob(
            user_id=user_id,
            original_filename=file_storage.filename,
            stored_filename=saved_name,
            total_rows=val_summary["total_rows"],
            valid_rows=val_summary["valid_rows"],
            positive_predictions=pos_count,
            negative_predictions=neg_count,
            status="completed",
            result_filepath=result_path,
            completed_at=datetime.now(timezone.utc),
        )
        db.session.add(job)
        db.session.flush()

        # Save individual result rows (up to 500 in DB for rapid pagination)
        for i, res in enumerate(batch_results[:500]):
            db_res = BulkPredictionResult(
                job_id=job.id,
                row_number=i + 1,
                inputs_json=json.dumps(df.iloc[i].to_dict()),
                prediction_class=res["prediction_class"],
                prediction_label=res["prediction_label"],
                probability=res["probability"],
            )
            db.session.add(db_res)

        db.session.commit()

        return {
            "job_id": job.id,
            "filename": file_storage.filename,
            "status": "completed",
            "total_records": job.total_rows,
            "positive_records": job.positive_predictions,
            "negative_records": job.negative_predictions,
            "results_sample": batch_results[:20],
            "download_url": f"/api/v1/bulk/jobs/{job.id}/download",
            "created_at": job.created_at.isoformat(),
        }

    @staticmethod
    def get_job_by_id(job_id: str, user_id: str = None) -> BulkJob:
        job = db.session.get(BulkJob, job_id)
        if not job:
            raise NotFoundError(f"Bulk job '{job_id}' not found.")
        if user_id and job.user_id and job.user_id != user_id:
            raise ForbiddenError("You do not have access to this bulk job.")
        return job

    @staticmethod
    def get_job_results_paginated(job_id: str, user_id: str = None, page: int = 1, per_page: int = 50) -> dict:
        job = BulkPredictionService.get_job_by_id(job_id, user_id)
        pagination = (
            BulkPredictionResult.query.filter_by(job_id=job.id)
            .order_by(BulkPredictionResult.row_number.asc())
            .paginate(page=page, per_page=per_page, error_out=False)
        )
        return {
            "job": job.to_dict(),
            "results": [item.to_dict() for item in pagination.items],
            "total": pagination.total,
            "pages": pagination.pages,
            "page": pagination.page,
            "per_page": pagination.per_page,
        }

    @staticmethod
    def generate_excel_template() -> io.BytesIO:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Patient Records"

        # Headers
        headers = REQUIRED_CLINICAL_FEATURES
        ws.append(headers)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill

        # Add sample row
        sample_row = [58, 1, 0, 140, 240, 0, 0, 150, 1, 1.5, 1, 0, 2]
        ws.append(sample_row)

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream
